"""Экспорт номенклатуры (nomenclature.export_all) — помимо основных полей
теперь выгружает и остаток по товару: сумма упакованного в короба (в
любом статусе/складе) и неразмещенного остатка (nomenclature._stock_by_item,
та же логика, что уже показывается в списке номенклатуры)."""

import io

import openpyxl

from wms.extensions import db
from wms.models import Box, BoxItem, Nomenclature, UnplacedStock, Warehouse


def _make_warehouse(code="WH-NOMEXP"):
    wh = Warehouse(code=code, name="Тест склад")
    db.session.add(wh)
    db.session.commit()
    return wh


def _read_rows(data):
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    return header, rows


def test_export_includes_stock_column_header(db, client_logged_in):
    item = Nomenclature(sku="SKU-NOMEXP-1", barcode="9992000001", name="Товар без остатка", unit="шт")
    db.session.add(item)
    db.session.commit()

    resp = client_logged_in.get("/nomenclature/export.xlsx")

    assert resp.status_code == 200
    header, rows = _read_rows(resp.data)
    assert header[-1] == "Остаток"
    row = next(r for r in rows if r[0] == item.barcode)
    assert row[-1] == 0


def test_export_sums_boxed_and_unplaced_stock(db, client_logged_in):
    wh = _make_warehouse()
    item = Nomenclature(sku="SKU-NOMEXP-2", barcode="9992000002", name="Товар с остатком", unit="шт")
    db.session.add(item)
    db.session.commit()

    box = Box(box_number="BOX-NOMEXP-1", warehouse_id=wh.id, status="open")
    db.session.add(box)
    db.session.commit()
    db.session.add(BoxItem(box_id=box.id, nomenclature_id=item.id, qty=5))
    db.session.add(UnplacedStock(warehouse_id=wh.id, nomenclature_id=item.id, qty=7))
    db.session.commit()

    resp = client_logged_in.get("/nomenclature/export.xlsx")

    _header, rows = _read_rows(resp.data)
    row = next(r for r in rows if r[0] == item.barcode)
    assert row[-1] == 12  # 5 в коробе + 7 неразмещенных


def test_export_ignores_zero_or_negative_unplaced_stock_rows(db, client_logged_in):
    wh = _make_warehouse("WH-NOMEXP-2")
    item = Nomenclature(sku="SKU-NOMEXP-3", barcode="9992000003", name="Товар с нулевым остатком", unit="шт")
    db.session.add(item)
    db.session.commit()
    db.session.add(UnplacedStock(warehouse_id=wh.id, nomenclature_id=item.id, qty=0))
    db.session.commit()

    resp = client_logged_in.get("/nomenclature/export.xlsx")

    _header, rows = _read_rows(resp.data)
    row = next(r for r in rows if r[0] == item.barcode)
    assert row[-1] == 0
