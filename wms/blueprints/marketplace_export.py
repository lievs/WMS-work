"""Выгрузка состава коробов перемещения в формате, который принимают
личные кабинеты маркетплейсов при заведении поставки:

- Ozon: "Состав грузовых мест" — короб WMS = одно грузовое место (ГМ).
  Ozon сам генерирует штрихкоды ГМ в своем кабинете (см. приложенный
  шаблон), поэтому их нужно вставить сюда самим, по одному на строку, в
  ТОМ ЖЕ порядке, что и короба в перемещении (см. movement.detail —
  порядок совпадает с MovementLine.id). "Артикул товара" Ozon — не наш
  внутренний SKU, а отдельная строка, которую нужно предварительно
  сопоставить со штрихкодом через OzonArticleMapping (см. ozon_mapping).
- Wildberries: тот же принцип, но проще — колонки "Баркод товара"/"ШК
  короба" не требуют отдельного сопоставления артикулов, годится наш
  собственный Nomenclature.barcode как есть.

И там, и там "Срок годности" WMS не отслеживает — оставляется пустым,
заполняется вручную при необходимости, как и предусмотрено самими
шаблонами."""

from flask import Blueprint, Response, abort, flash, redirect, render_template, request, url_for
from flask_login import current_user
from openpyxl import load_workbook

from ..extensions import db
from ..models import MovementDocument, MovementLine, OzonArticleMapping
from ..utils.excel_io import (
    export_ozon_package_composition,
    export_wb_package_composition,
    timestamp_for_filename,
)
from ..utils.http import content_disposition
from .movement import _can_view_movement_document

bp = Blueprint("marketplace_export", __name__)


def _get_viewable_movement(doc_id):
    doc = MovementDocument.query.get_or_404(doc_id)
    if not _can_view_movement_document(doc):
        abort(404)
    return doc


def _cell_to_barcode(value):
    """Штрихкод в файле сопоставления приходит из Excel как число (Ozon
    отдает баркоды без ведущих нулей и без научной нотации) — приводим к
    обычной строке без ".0" на конце, а не str(float(...))."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


@bp.route("/ozon-mapping", methods=["GET", "POST"])
def ozon_mapping():
    if not current_user.is_admin:
        flash("Загружать сопоставление артикулов Ozon может только администратор", "danger")
        return redirect(url_for("main.index"))

    if request.method == "POST":
        file = request.files.get("file")
        if not file or file.filename == "":
            flash("Выберите файл xlsx", "danger")
            return redirect(url_for("marketplace_export.ozon_mapping"))

        workbook = load_workbook(file, read_only=True, data_only=True)
        worksheet = workbook.active
        updated = 0
        for row in worksheet.iter_rows(values_only=True):
            if not row or row[0] is None:
                continue
            barcode = _cell_to_barcode(row[0])
            article = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            if not barcode or not article:
                continue
            mapping = OzonArticleMapping.query.filter_by(barcode=barcode).first()
            if mapping is None:
                mapping = OzonArticleMapping(barcode=barcode)
                db.session.add(mapping)
            mapping.article = article
            updated += 1
        db.session.commit()
        flash(f"Сопоставление артикулов Ozon обновлено: {updated} строк", "success")
        return redirect(url_for("marketplace_export.ozon_mapping"))

    return render_template(
        "marketplace_export/ozon_mapping.html", count=OzonArticleMapping.query.count()
    )


@bp.route("/movement/<int:doc_id>/ozon", methods=["GET", "POST"])
def ozon_package_composition(doc_id):
    doc = _get_viewable_movement(doc_id)
    lines = doc.lines.order_by(MovementLine.id.asc()).all()

    if request.method == "POST":
        raw = request.form.get("cargo_barcodes", "")
        cargo_barcodes = [b.strip() for b in raw.splitlines() if b.strip()]
        if len(cargo_barcodes) != len(lines):
            flash(
                f"Штрихкодов ГМ должно быть ровно {len(lines)} (по числу коробов в "
                f"перемещении), а введено {len(cargo_barcodes)}",
                "danger",
            )
            return render_template(
                "marketplace_export/ozon_export.html", doc=doc, lines=lines, cargo_barcodes_raw=raw
            )

        rows = []
        unmapped = set()
        for line, cargo_barcode in zip(lines, cargo_barcodes):
            for item in line.box.items:
                barcode = item.nomenclature.barcode
                mapping = OzonArticleMapping.query.filter_by(barcode=barcode).first()
                if mapping is None:
                    unmapped.add(barcode)
                rows.append(
                    {
                        "barcode": barcode,
                        "article": mapping.article if mapping else "",
                        "qty": item.qty,
                        "cargo_barcode": cargo_barcode,
                    }
                )

        data = export_ozon_package_composition(rows)
        if unmapped:
            flash(
                "Не найден артикул Ozon для штрихкодов: " + ", ".join(sorted(unmapped))
                + " — колонка «Артикул товара» для них оставлена пустой, заполните вручную "
                "или дозагрузите сопоставление.",
                "warning",
            )
        fname = f"{doc.number}_ozon_gm_{timestamp_for_filename()}.xlsx"
        return Response(
            data,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": content_disposition(fname)},
        )

    return render_template(
        "marketplace_export/ozon_export.html", doc=doc, lines=lines, cargo_barcodes_raw=""
    )


@bp.route("/movement/<int:doc_id>/wb", methods=["GET", "POST"])
def wb_package_composition(doc_id):
    doc = _get_viewable_movement(doc_id)
    lines = doc.lines.order_by(MovementLine.id.asc()).all()

    if request.method == "POST":
        raw = request.form.get("box_barcodes", "")
        box_barcodes = [b.strip() for b in raw.splitlines() if b.strip()]
        if len(box_barcodes) != len(lines):
            flash(
                f"Штрихкодов короба должно быть ровно {len(lines)} (по числу коробов в "
                f"перемещении), а введено {len(box_barcodes)}",
                "danger",
            )
            return render_template(
                "marketplace_export/wb_export.html", doc=doc, lines=lines, box_barcodes_raw=raw
            )

        rows = []
        for line, box_barcode in zip(lines, box_barcodes):
            for item in line.box.items:
                rows.append(
                    {
                        "barcode": item.nomenclature.barcode,
                        "qty": item.qty,
                        "box_barcode": box_barcode,
                    }
                )

        data = export_wb_package_composition(rows)
        fname = f"{doc.number}_wb_shk_{timestamp_for_filename()}.xlsx"
        return Response(
            data,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": content_disposition(fname)},
        )

    return render_template(
        "marketplace_export/wb_export.html", doc=doc, lines=lines, box_barcodes_raw=""
    )
